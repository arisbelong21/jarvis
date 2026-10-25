"""
plugins/docx_writer.py — makes a real .docx from a spoken request.

The companion of excel_writer: a spoken "buatkan dokumen Word tentang …"
becomes an actual Word file with a title, headings, paragraphs and bullets.
Content is written by the heavy-task ladder (custom provider first), so the
document is as good as the model doing the writing; the file itself is built
locally with python-docx — no cloud document service involved.

Where it saves follows the same rule as every other writer (core/workspace):
the folder the user named → the last project folder → the Desktop. That is
what keeps "taruh di folder kerja saya" true.
"""

import re
from datetime import datetime
from pathlib import Path

PLUGIN = {
    "name": "docx_writer",
    "description": (
        "Creates a real Word .docx document FILE from a spoken description: "
        "title, headings, paragraphs, bullet lists. Use for 'buatkan dokumen "
        "Word tentang…', 'tulis laporan … dalam Word', 'make me a docx about…', "
        "'susun proposal dalam Word'. Pass the user's full request verbatim in "
        "'request'. Use this only when the user wants an actual Word FILE — "
        "not for answering a question or for a chat reply."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "request": {
                "type": "STRING",
                "description": ("The user's full document request, verbatim, "
                                "in their own language."),
            },
            "filename": {
                "type": "STRING",
                "description": ("Optional file name (without extension). If "
                                "omitted, one is derived from the request."),
            },
            "folder": {
                "type": "STRING",
                "description": (
                    "Folder tujuan (path lengkap) kalau pengguna menyebut "
                    "lokasi tertentu atau 'folder kerja saya'. Kosongkan untuk "
                    "memakai folder kerja terakhir, lalu Desktop."
                ),
            },
        },
        "required": ["request"],
    },
}


def _safe_name(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*]', "", (name or "").strip()) or "dokumen"
    return name[:60]


def _output_dir(explicit: str = "") -> Path:
    try:
        from core import workspace
        return workspace.target_dir(explicit)
    except Exception:
        pass
    if explicit:
        try:
            return Path(explicit).expanduser()
        except Exception:
            pass
    desktop = Path.home() / "Desktop"
    return desktop if desktop.is_dir() else Path.home()


def _build_spec(request: str) -> dict:
    """Ask the ladder for a structured outline — JSON only, no prose."""
    from core import gemini

    prompt = (
        "Susun kerangka dokumen Word dari permintaan berikut. Balas HANYA JSON "
        "minified, tanpa markdown:\n"
        '{"filename": "nama_file_tanpa_ekstensi", "title": "Judul dokumen", '
        '"sections": [{"heading": "Judul bagian", "paragraphs": ["paragraf '
        'utuh"], "bullets": ["poin singkat"]}]}\n'
        "Aturan: 3–8 bagian; paragraf harus berisi kalimat lengkap yang siap "
        "dibaca (bukan placeholder); gunakan bullets hanya bila memang daftar; "
        "bahasa mengikuti permintaan pengguna.\n"
        f"PERMINTAAN: {request}"
    )
    spec = gemini.as_json(prompt, tier=gemini.SMART, timeout_ms=120000,
                          default=None)
    if not isinstance(spec, dict):
        raise RuntimeError("model tidak mengembalikan kerangka yang valid")
    if not spec.get("sections"):
        raise RuntimeError("kerangka kosong")
    return spec


def _build_docx(spec: dict, path: Path) -> dict:
    from docx import Document

    doc = Document()
    title = str(spec.get("title") or "Dokumen").strip()
    doc.add_heading(title, level=0)

    sections = spec.get("sections") or []
    paragraphs = bullets = 0
    for sec in sections:
        if not isinstance(sec, dict):
            continue
        heading = str(sec.get("heading") or "").strip()
        if heading:
            doc.add_heading(heading, level=1)
        for text in sec.get("paragraphs") or []:
            text = str(text or "").strip()
            if text:
                doc.add_paragraph(text)
                paragraphs += 1
        for text in sec.get("bullets") or []:
            text = str(text or "").strip()
            if text:
                try:
                    doc.add_paragraph(text, style="List Bullet")
                except Exception:
                    doc.add_paragraph("• " + text)
                bullets += 1

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return {"sections": len(sections), "paragraphs": paragraphs,
            "bullets": bullets, "title": title}


def run(parameters: dict, player=None, session_memory=None) -> str:
    request = (parameters.get("request") or "").strip()
    if not request:
        return "Ceritakan isi dokumen yang bos inginkan."

    try:
        import docx  # noqa: F401
    except ImportError:
        return ("Saya butuh pustaka python-docx untuk membuat file Word. "
                "Pasang dengan: pip install python-docx")

    try:
        spec = _build_spec(request)
    except Exception as e:                           # noqa: BLE001
        return f"Kerangka dokumennya gagal disusun: {e}"

    fname = _safe_name(parameters.get("filename")
                       or spec.get("filename") or "dokumen")
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out_dir = _output_dir(str(parameters.get("folder") or "").strip())
    try:
        out_dir.mkdir(parents=True, exist_ok=True)
    except Exception:
        out_dir = _output_dir()
    path = out_dir / f"{fname}_{stamp}.docx"

    try:
        info = _build_docx(spec, path)
    except Exception as e:                           # noqa: BLE001
        return f"Kerangkanya siap, tapi file Word gagal ditulis: {e}"

    if player:
        try:
            player.show_content(
                "📄 WORD CREATED",
                f"{path.name}\n\n{info['sections']} bagian · "
                f"{info['paragraphs']} paragraf · {info['bullets']} poin\n\n"
                f"Saved to:\n{path}")
        except Exception:
            pass

    return (f"Selesai — {path.name} sudah dibuat "
            f"({info['sections']} bagian, {info['paragraphs']} paragraf) "
            f"dan disimpan di {path.parent}.")
